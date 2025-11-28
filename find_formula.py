import openpyxl
import math

wb = openpyxl.load_workbook('data/find_spp27_11_25_1.xlsx', read_only=True)
ws = wb.active

print('Анализ формул для столбцов T и U:\n')

# Собираем данные из всех строк
data_rows = []
for row in range(3, min(20, ws.max_row + 1)):
    try:
        o = ws.cell(row=row, column=15).value
        p = ws.cell(row=row, column=16).value
        s = ws.cell(row=row, column=19).value
        t = ws.cell(row=row, column=20).value
        u = ws.cell(row=row, column=21).value
        w = ws.cell(row=row, column=23).value
        r = ws.cell(row=row, column=18).value
        q = ws.cell(row=row, column=17).value
        
        if all(isinstance(v, (int, float)) for v in [o, p, s, t, u] if v is not None):
            data_rows.append({
                'row': row,
                'O': float(o),
                'P': float(p),
                'S': float(s),
                'T': float(t),
                'U': float(u),
                'W': float(w) if w and isinstance(w, (int, float)) else None,
                'R': float(r) if r and isinstance(r, (int, float)) else None,
                'Q': float(q) if q and isinstance(q, (int, float)) else None,
            })
    except (ValueError, TypeError):
        continue

print(f'Найдено {len(data_rows)} строк с числовыми данными\n')

# Проверяем различные формулы
formulas_to_test = [
    # Простые формулы
    ('U = O - S', lambda d: d['O'] - d['S']),
    ('U = O - P', lambda d: d['O'] - d['P']),
    ('U = O - W', lambda d: d['O'] - d['W'] if d['W'] else None),
    
    # Комбинации
    ('U = (O - S) + (O - P)', lambda d: (d['O'] - d['S']) + (d['O'] - d['P'])),
    ('U = (O - S) + R', lambda d: (d['O'] - d['S']) + d['R'] if d['R'] else None),
    ('U = O - S + R', lambda d: d['O'] - d['S'] + d['R'] if d['R'] else None),
    
    # С процентами
    ('U = O * (T/100)', lambda d: d['O'] * (d['T'] / 100)),
    ('U = O * (1 - S/O)', lambda d: d['O'] * (1 - d['S'] / d['O'])),
    
    # Обратные расчеты
    ('U = O - (O * S / P)', lambda d: d['O'] - (d['O'] * d['S'] / d['P']) if d['P'] else None),
    ('U = O - S * (O/P)', lambda d: d['O'] - d['S'] * (d['O'] / d['P']) if d['P'] else None),
    
    # С минимальной ценой
    ('U = O - W + (O - S)', lambda d: (d['O'] - d['W']) + (d['O'] - d['S']) if d['W'] else None),
    ('U = (O - W) + (O - S)', lambda d: (d['O'] - d['W']) + (d['O'] - d['S']) if d['W'] else None),
    
    # Проверка через T
    ('U = O * (T/100) (из эталона)', lambda d: d['O'] * (d['T'] / 100)),
    
    # Новые варианты
    ('U = O - S + (O - P)', lambda d: d['O'] - d['S'] + (d['O'] - d['P'])),
    ('U = 2*O - S - P', lambda d: 2 * d['O'] - d['S'] - d['P']),
]

print('Проверка формул для U:\n')
for formula_name, formula_func in formulas_to_test:
    matches = 0
    total_diff = 0
    max_diff = 0
    
    for d in data_rows:
        try:
            calculated = formula_func(d)
            if calculated is None:
                continue
            diff = abs(calculated - d['U'])
            total_diff += diff
            max_diff = max(max_diff, diff)
            if diff < 0.1:  # Точное совпадение
                matches += 1
        except (ZeroDivisionError, TypeError):
            continue
    
    if len(data_rows) > 0:
        avg_diff = total_diff / len(data_rows)
        print(f'{formula_name}:')
        print(f'  Точных совпадений: {matches}/{len(data_rows)}')
        print(f'  Средняя разница: {avg_diff:.2f}')
        print(f'  Максимальная разница: {max_diff:.2f}')
        
        # Показываем примеры для первых 3 строк
        for i, d in enumerate(data_rows[:3]):
            try:
                calc = formula_func(d)
                if calc is not None:
                    print(f'    Строка {d["row"]}: {calc:.2f} (эталон: {d["U"]}, разница: {abs(calc - d["U"]):.2f})')
            except:
                pass
        print()

# Обратная логика: если T = round((U/O)*100), то откуда U?
print('\nОбратная логика: вычисляем U, затем T = round((U/O)*100):\n')

# Пробуем различные способы вычисления U, затем проверяем T
u_candidates = [
    ('U = O - S', lambda d: d['O'] - d['S']),
    ('U = (O - S) + (O - P)', lambda d: (d['O'] - d['S']) + (d['O'] - d['P'])),
    ('U = O - S + (O - P)', lambda d: d['O'] - d['S'] + (d['O'] - d['P'])),
    ('U = 2*O - S - P', lambda d: 2 * d['O'] - d['S'] - d['P']),
    ('U = O - S + R', lambda d: d['O'] - d['S'] + d['R'] if d['R'] else None),
    ('U = O * (1 - S/O) + (O - P)', lambda d: d['O'] * (1 - d['S']/d['O']) + (d['O'] - d['P'])),
    ('U = (O - S) * (O / P)', lambda d: (d['O'] - d['S']) * (d['O'] / d['P']) if d['P'] else None),
    ('U = O - (S * P / O)', lambda d: d['O'] - (d['S'] * d['P'] / d['O']) if d['P'] else None),
    ('U = round(O - (S * P / O))', lambda d: round(d['O'] - (d['S'] * d['P'] / d['O'])) if d['P'] else None),
    ('U = O - round(S * P / O)', lambda d: d['O'] - round(d['S'] * d['P'] / d['O']) if d['P'] else None),
    ('U = round((O - S) * (O / P))', lambda d: round((d['O'] - d['S']) * (d['O'] / d['P'])) if d['P'] else None),
    ('U = (O - S) + round((O - P) * S / P)', lambda d: (d['O'] - d['S']) + round((d['O'] - d['P']) * d['S'] / d['P']) if d['P'] else None),
    ('U = O - round(S * (1 - P/O))', lambda d: d['O'] - round(d['S'] * (1 - d['P']/d['O'])) if d['P'] else None),
]

for u_name, u_func in u_candidates:
    matches_t = 0
    matches_u = 0
    total_diff_t = 0
    total_diff_u = 0
    
    for d in data_rows:
        try:
            u_calc = u_func(d)
            if u_calc is None:
                continue
            
            # Вычисляем T из вычисленного U
            t_calc = round((u_calc / d['O']) * 100)
            
            diff_t = abs(t_calc - d['T'])
            diff_u = abs(u_calc - d['U'])
            
            total_diff_t += diff_t
            total_diff_u += diff_u
            
            if diff_t < 0.1:
                matches_t += 1
            if diff_u < 0.1:
                matches_u += 1
        except (ZeroDivisionError, TypeError):
            continue
    
    if len(data_rows) > 0:
        avg_diff_t = total_diff_t / len(data_rows)
        avg_diff_u = total_diff_u / len(data_rows)
        print(f'{u_name}, затем T = round((U/O)*100):')
        print(f'  Точных совпадений T: {matches_t}/{len(data_rows)}, средняя разница: {avg_diff_t:.2f}')
        print(f'  Точных совпадений U: {matches_u}/{len(data_rows)}, средняя разница: {avg_diff_u:.2f}')
        
        # Показываем примеры
        for i, d in enumerate(data_rows[:3]):
            try:
                u_calc = u_func(d)
                if u_calc is not None:
                    t_calc = round((u_calc / d['O']) * 100)
                    print(f'    Строка {d["row"]}: U={u_calc:.2f} (эталон: {d["U"]}), T={t_calc}% (эталон: {d["T"]}%)')
            except:
                pass
        print()

wb.close()

